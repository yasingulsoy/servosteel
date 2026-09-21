# -*- coding: utf-8 -*-
"""seo/hedef-firmalar/*.md icindeki dogrulanmis firma tablolarini tek Excel'e cevirir.

  python scripts/hedef-firma-excel.py

Cikti: seo/hedef-firmalar/Servosteel-Hedef-Firmalar.xlsx

Her segment bir sayfa; bastaki "TUM FIRMALAR" sayfasi hepsini birlestirir.
Calisma dosyasi oldugu icin her satira Durum / Gonderim / Yanit / Not sutunlari
eklenir — liste okunacak degil, uzerinde calisilacak.

Tablolar farkli ajanlardan geldigi icin BASLIKLAR NORMALLESTIRILIR: "Company",
"Firma", "Company Name" hepsi tek sutuna oturur. Taninmayan sutun atilmaz,
sona eklenir — veri kaybetmektense fazladan sutun iyidir.
"""
import io
import os
import re
import sys
import unicodedata
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KLASOR = os.path.join(KOK, "seo", "hedef-firmalar")
CIKTI = os.path.join(KLASOR, "Servosteel-Hedef-Firmalar.xlsx")

ADLAR = {
    "kablo-kanali": "Kablo Kanalı",
    "solar-profil": "Solar Profil",
    "celik-servis-merkezi": "Çelik Servis Merkezi",
    "yol-bariyeri": "Yol Bariyeri",
    "raf-sistemleri": "Raf Sistemleri",
    "alcipan-profili": "Alçıpan Profili",
    "cati-cephe-paneli": "Çatı ve Cephe Paneli",
    "pres-atolyeleri": "Pres Atölyeleri",
    "asik-celik-yapi": "Aşık ve Çelik Yapı",
    "iskele-kalasi": "İskele Kalası",
    "market-rafi": "Market Rafı",
    "gurultu-bariyeri": "Gürültü Bariyeri",
    "havalandirma-kanali": "Havalandırma Kanalı",
    "metal-mobilya": "Çelik Mobilya",
}

# Kanonik sutunlar ve onlara isaret eden kaliplar (once eslesen kazanir)
SEMA = [
    ("Firma",              r"^(firma|company|company name|şirket|sirket|name)$"),
    ("Ülke",               r"^(ülke|ulke|country|location)$"),
    ("Web sitesi",         r"^(web sitesi|website|web site|site|url|web)$"),
    ("Doğrulama sayfası",  r"(doğrulama|dogrulama|verification|evidence|kanıt|kanit|proof)"),
    ("Ne üretiyor",        r"(ne üretiyor|ne uretiyor|what|ürün|urun|product|note|not|açıklama|aciklama|size|büyüklük)"),
    ("E-posta / İletişim", r"(e-?posta|e-?mail|mail|iletişim|iletisim|contact)"),
]
TAKIP = ["Durum", "Gönderim tarihi", "Yanıt", "Not"]
LINK_BASLIK = "Gönderilecek link"

# Sitede ve katalogda karsiligi olmayan segment. Firma "bu hatti uretiyoruz"
# demeden bu firmalara yazilmasin diye Durum sutunu onceden doldurulur.
TEYITSIZ = {}   # 2026-09-21: alcipan sorusu cevaplandi (hat yok, besleyici teklif edilir)

# Cok segmentli firmada e-posta HAT teklifiyle acilir; yalnizca besleyici/acici
# teklif edilen segment sona gider (Dana Steel'e once trapez hatti, besleyici en son).
SONA = {"Alçıpan Profili", "Pres Atölyeleri", "Havalandırma Kanalı", "Çelik Mobilya"}


def durum_isaretle(ws, segment_sutunu=None, segment_adi=None):
    """Teyitsiz segmentteki satirlarin Durum hucresini doldurur."""
    bas = [c.value for c in ws[1]]
    if "Durum" not in bas:
        return 0
    di = bas.index("Durum") + 1
    si = bas.index("Segment") + 1 if "Segment" in bas else None
    n = 0
    for r in range(2, ws.max_row + 1):
        seg = str(ws.cell(row=r, column=si).value) if si else segment_adi
        if seg in TEYITSIZ:          # yalnizca TEK segment teyitsizse; cok segmentli firma gonderilebilir
            ws.cell(row=r, column=di).value = TEYITSIZ[seg]
            ws.cell(row=r, column=di).font = Font(bold=True, color="9C0006")
            n += 1
    return n

# Ulke sutunu suzgecle kullanilacak; ajanlar "BAE", "Birleşik Arap Emirlikleri",
# "BAE (Dubai/Ajman)" diye uc ayri deger yaziyor ve suzgec bolunuyor.
# Once parantezli sehir bilgisi atilir, sonra ad tek yazima cekilir.
ULKE_ESI = {
    "bae": "BAE", "birleşik arap emirlikleri": "BAE", "uae": "BAE",
    "birlesik arap emirlikleri": "BAE", "united arab emirates": "BAE",
    "suudi arabistan": "Suudi Arabistan", "saudi arabia": "Suudi Arabistan",
    "hindistan": "Hindistan", "india": "Hindistan",
    "polonya": "Polonya", "poland": "Polonya",
    "italya": "İtalya", "i̇talya": "İtalya", "italy": "İtalya",
    "ispanya": "İspanya", "i̇spanya": "İspanya", "spain": "İspanya",
    "mısır": "Mısır", "misir": "Mısır", "egypt": "Mısır",
    "fas": "Fas", "morocco": "Fas", "maroc": "Fas",
    "birleşik krallık": "Birleşik Krallık", "birlesik krallik": "Birleşik Krallık",
    "uk": "Birleşik Krallık", "united kingdom": "Birleşik Krallık",
    "ingiltere": "Birleşik Krallık", "i̇ngiltere": "Birleşik Krallık",
    "abd": "ABD", "usa": "ABD", "amerika": "ABD", "united states": "ABD",
    "almanya": "Almanya", "germany": "Almanya",
    "romanya": "Romanya", "romania": "Romanya",
    "brezilya": "Brezilya", "brazil": "Brezilya",
    "kolombiya": "Kolombiya", "colombia": "Kolombiya",
    "meksika": "Meksika", "mexico": "Meksika",
    "güney afrika": "Güney Afrika", "guney afrika": "Güney Afrika",
    "south africa": "Güney Afrika",
    "endonezya": "Endonezya", "indonesia": "Endonezya",
    "kenya": "Kenya", "vietnam": "Vietnam", "litvanya": "Litvanya",
    "lithuania": "Litvanya", "özbekistan": "Özbekistan", "ozbekistan": "Özbekistan",
    "uzbekistan": "Özbekistan", "türkiye": "Türkiye", "turkiye": "Türkiye",
    "turkey": "Türkiye", "nijerya": "Nijerya", "nigeria": "Nijerya",
    "kazakistan": "Kazakistan", "kazakhstan": "Kazakistan",
    "çekya": "Çekya", "cekya": "Çekya", "czechia": "Çekya",
    "portekiz": "Portekiz", "portugal": "Portekiz",
    "malezya": "Malezya", "malaysia": "Malezya",
    "tayland": "Tayland", "thailand": "Tayland",
    "filipinler": "Filipinler", "philippines": "Filipinler",
    "katar": "Katar", "qatar": "Katar", "irak": "Irak", "iraq": "Irak",
    "peru": "Peru", "şili": "Şili", "sili": "Şili", "chile": "Şili",
    "gana": "Gana", "ghana": "Gana", "tanzanya": "Tanzanya",
    "azerbaycan": "Azerbaycan", "gürcistan": "Gürcistan",
}


# Gonderilecek link: segmentin sayfasi, firmanin ulkesinin dilinde, UTM'li.
# UTM sayesinde GA4'te "outreach / email" kaynagi segment ve firma kiriliminda
# gorunur — hangi firma linke tikladi, takip telefonu ona gore acilir.
SITE = "https://servosteel.com.tr"
SEGMENT_SAYFA = {   # segment -> (TR yol, diger diller yolu, kampanya adi)
    "Kablo Kanalı":         ("/roll-form-hatlari/kablo-kanali", "/roll-forming-lines/cable-tray", "kablo-kanali"),
    "Solar Profil":         ("/roll-form-hatlari/solar-profil", "/roll-forming-lines/solar-panel-profile", "solar-profil"),
    "Raf Sistemleri":       ("/roll-form-hatlari/agir-raf", "/roll-forming-lines/storage-rack", "raf"),
    "Yol Bariyeri":         ("/roll-form-hatlari/yol-bariyeri", "/roll-forming-lines/guard-rail", "yol-bariyeri"),
    "Çatı ve Cephe Paneli": ("/roll-form-hatlari/trapez-cephe-paneli", "/roll-forming-lines/trapezoidal-and-facade-panel", "cati-panel"),
    "Çelik Servis Merkezi": ("/dilme-hatlari", "/coil-slitting-lines", "celik-servis"),
    "Pres Atölyeleri":      ("/makineler/servo-suruculer", "/machines/servo-feeders", "pres-atolyesi"),
    # Firma alcipan HATTI uretmiyor (katalogda yok, kullanici 2026-09-21). Bu firmalara
    # katalogdaki acici + dogrultmali servo surucu teklif edilir; kanit: firmanin
    # "Asma Tavan Hattı / Servo Sürücü Ve Rulo Sac Açma Sistemleri" videosu (2017).
    "Alçıpan Profili":      ("/makineler/dogrultmali-servo-suruculer", "/machines/straightener-servo-feeders", "alcipan"),
    "Aşık ve Çelik Yapı":   ("/roll-form-hatlari/c-sigma-omega", "/roll-forming-lines/c-sigma-omega-profiles", "asik"),
    "İskele Kalası":        ("/roll-form-hatlari/iskele-kalas", "/roll-forming-lines/scaffolding-plank", "iskele"),
    # Market rafi: sektor sayfasi "magaza rafi profilleri + teshir elemanlari"ni anlatiyor.
    "Market Rafı":          ("/uygulamalar/metal-mobilya", "/applications/metal-furniture-and-retail", "market-rafi"),
    "Gürültü Bariyeri":     ("/roll-form-hatlari/gurultu-bariyeri", "/roll-forming-lines/noise-barrier", "gurultu"),
    # Kanal hattinin girisi = acici + dogrultucu + besleyici; kompakt hat tam bu.
    "Havalandırma Kanalı":  ("/makineler/kompakt-hatlar", "/machines/compact-lines", "havalandirma"),
    "Çelik Mobilya":        ("/uygulamalar/metal-mobilya", "/applications/metal-furniture-and-retail", "mobilya"),
}
ULKE_DIL = {
    "İspanya": "es", "Meksika": "es", "Kolombiya": "es", "Peru": "es", "Şili": "es",
    "Polonya": "pl", "İtalya": "it", "Almanya": "de", "Macaristan": "hu",
    "Kazakistan": "ru", "Özbekistan": "ru", "Kırgızistan": "ru", "Tacikistan": "ru", "Türkiye": "tr",
    # E-posta dili Fransizca / Portekizce; sitede bu diller yok, link /en/ sayfasina gider
    "Fransa": "fr", "Belçika": "fr", "Lüksemburg": "fr", "Fas": "fr", "Cezayir": "fr", "Tunus": "fr",
    "Senegal": "fr", "Fildişi Sahili": "fr", "Kamerun": "fr", "Mali": "fr", "Burkina Faso": "fr",
    "Benin": "fr", "Togo": "fr", "Gabon": "fr", "Kongo DC": "fr",
    "Brezilya": "pt", "Portekiz": "pt", "Angola": "pt", "Mozambik": "pt",
    "Ekvador": "es", "Bolivya": "es", "Paraguay": "es", "Uruguay": "es", "Arjantin": "es",
    "Guatemala": "es", "Honduras": "es", "El Salvador": "es", "Kosta Rika": "es", "Panama": "es",
    "Dominik Cumhuriyeti": "es",
}
SITE_DILLERI = {"tr", "en", "de", "es", "it", "hu", "pl", "ru", "ar"}  # geri kalan her ulke: en (Korfez ve Kuzey Afrika is dunyasi Ingilizce yazisiyor)


def gonderim_linki(segment, ulke, satir):
    tr_yol, en_yol, kampanya = SEGMENT_SAYFA.get(segment, ("/", "/", "genel"))
    dil = ULKE_DIL.get(ulke, "en")
    if dil not in SITE_DILLERI:      # fr/pt: e-posta o dilde, sayfa Ingilizce (sitede yok)
        dil = "en"
    yol = tr_yol if dil == "tr" else "/%s%s" % (dil, en_yol)
    m = re.search(r"https?://(?:www\.)?([^/\s|]+)", str(satir[2] or ""))
    icerik = re.sub(r"[^a-z0-9.-]", "", (m.group(1) if m else sade(satir[0])).lower())[:40]
    return "%s%s?utm_source=outreach&utm_medium=email&utm_campaign=%s&utm_content=%s" % (
        SITE, yol, kampanya, icerik)


# Hazir e-posta: seo/eposta-taslaklari.json (iddialarin hepsi sitede yaziyor).
import json
from urllib.parse import quote
SABLON = json.load(io.open(os.path.join(KOK, "seo", "eposta-taslaklari.json"), encoding="utf-8"))
EK_BASLIK = [LINK_BASLIK, "Konu", "Hazır e-posta", "Taslak"]
MAILTO_SINIR = 2000   # Outlook ve Excel koprusu ~2080 karakterde kesiyor


def hitap_adi(firma):
    """Selamlamada kullanilacak kisa ad: parantez ve sirket turu eki atilir."""
    a = re.sub(r"\(.*?\)|[«»\"“”*]", " ", str(firma or ""))
    ek = (r"(,?\s*(pvt\.?\s*ltd\.?|private limited|ltd\.?|llc|l\.l\.c\.?|inc\.?|s\.a\.?\s*de\s*c\.v\.?|"
          r"s\.a\.?|s\.l\.?|s\.r\.l\.?|s\.r\.o\.?|sp\.?\s*z\s*o\.?\s*o\.?|gmbh|jsc|co\.,?|company limited|"
          r"s\.a\.r\.l\.?|plc|sdn\.?\s*bhd\.?|bhd\.?|pty\.?\s*ltd\.?|s\.p\.a\.?|spa|ltda\.?|a\.s\.?|jsc))\s*$")
    for _ in range(3):
        a = re.sub(ek, "", a.strip(" ,.-"), flags=re.I)
    # Bastaki sirket turu: Endonezya PT/CV, Kazakistan/Rusya TOO/OOO
    a = re.sub(r"^\s*(pt\.?|cv\.?|too|тоо|ooo|ооо)\s+", "", a, flags=re.I)
    return re.sub(r"\s+", " ", a).strip(" ,.-") or str(firma)


def ek_sutunlar(segler, satir):
    """[link, konu, metin, taslak]. segler: teyitli segment basta."""
    seg = segler[0]
    link = gonderim_linki(seg, satir[1], satir)
    if seg in TEYITSIZ:
        return [link, "", "", ""]
    dil = ULKE_DIL.get(satir[1], "en")
    t = SABLON.get(dil) or SABLON["en"]
    kamp = SEGMENT_SAYFA[seg][2]
    firma = hitap_adi(satir[0])
    digerleri = [SEGMENT_SAYFA[x][2] for x in segler[1:] if x not in TEYITSIZ]
    ek = ""
    if digerleri:
        adlar = [t["segment"][k]["ad"] for k in digerleri]
        liste = adlar[0] if len(adlar) == 1 else ", ".join(adlar[:-1]) + t["ve"] + adlar[-1]
        ek = t["ek"].replace("{liste}", liste)
    metin = (t["govde"].replace("{selam}", t["selam"]).replace("{firma}", firma)
             .replace("{cumle}", t["segment"][kamp]["cumle"]).replace("{ek}", ek)
             .replace("{link}", link))
    konu = t["segment"][kamp]["konu"]
    eposta = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", str(satir[5] or ""))
    taslak = ""
    if eposta:
        m = "mailto:%s?subject=%s&body=%s" % (eposta.group(0), quote(konu), quote(metin))
        if len(m) <= MAILTO_SINIR:
            taslak = m
    return [link, konu, metin, taslak]


# Ayni firmanin iki alan adi. Kanit: iki alan adindaki satirda da ayni tuzel kisi
# ("Dana Steel Processing Industry LLC") yaziyor.
ALAN_ESI = {"danagroups.com": "danasteeluae.com"}


SEGMENT_ESI = {   # bolge ajanlarinin yazabilecegi Ingilizce karsiliklar
    "cable tray": "Kablo Kanalı", "solar": "Solar Profil", "solar mounting": "Solar Profil",
    "racking": "Raf Sistemleri", "storage racking": "Raf Sistemleri", "pallet racking": "Raf Sistemleri",
    "guardrail": "Yol Bariyeri", "road barrier": "Yol Bariyeri",
    "roofing": "Çatı ve Cephe Paneli", "roofing and cladding": "Çatı ve Cephe Paneli", "sandwich panel": "Çatı ve Cephe Paneli",
    "steel service centre": "Çelik Servis Merkezi", "steel service center": "Çelik Servis Merkezi", "coil processing": "Çelik Servis Merkezi",
    "press shop": "Pres Atölyeleri", "metal stamping": "Pres Atölyeleri", "stamping": "Pres Atölyeleri",
    "drywall": "Alçıpan Profili", "drywall profile": "Alçıpan Profili", "ceiling profile": "Alçıpan Profili",
    "purlin": "Aşık ve Çelik Yapı", "peb": "Aşık ve Çelik Yapı", "steel building": "Aşık ve Çelik Yapı",
    "scaffolding": "İskele Kalası", "scaffolding plank": "İskele Kalası",
    "retail shelving": "Market Rafı", "gondola shelving": "Market Rafı",
    "noise barrier": "Gürültü Bariyeri", "hvac": "Havalandırma Kanalı", "ductwork": "Havalandırma Kanalı",
    "ventilation": "Havalandırma Kanalı", "steel furniture": "Çelik Mobilya", "metal furniture": "Çelik Mobilya",
}


def segment_bul(ham):
    """Ajanin yazdigi segment metnini 14 kanonik segmentten birine esler; bulamazsa ''."""
    d = sade(re.sub(r"\(.*?\)", " ", str(ham or ""))).strip(" .")
    for ad in ADLAR.values():
        if d == sade(ad):
            return ad
    for anahtar, ad in ADLAR.items():
        if d == anahtar or d == anahtar.replace("-", " "):
            return ad
    if d in SEGMENT_ESI:
        return SEGMENT_ESI[d]
    for ad in ADLAR.values():                 # "Kablo Kanalı + Raf" gibi: ilk tanınanı al
        if sade(ad) in d:
            return ad
    return ""


def firma_anahtari(satir):
    """Tekillestirme anahtari: alan adi + ulke.

    Ad yetmez — "PRK Steel" ile "PRK Steel Products" ayni firma, iki satir
    kaliyordu. Alan adi da tek basina yetmez — NS BlueScope Endonezya ve
    Vietnam ayni alan adinda ama ayri tesis, ayri muhatap."""
    m = re.search(r"https?://(?:www\.)?([^/\s|]+)", str(satir[2] or "")) \
        or re.search(r"https?://(?:www\.)?([^/\s|]+)", str(satir[3] or ""))
    if not m:
        return "ad:" + sade(satir[0]) + "|" + sade(satir[1])
    alan = m.group(1).lower().rstrip(".")
    return ALAN_ESI.get(alan, alan) + "|" + sade(satir[1])


def ulke_duzelt(ham):
    """'BAE (Dubai/Ajman)' -> 'BAE'. Coklu ulke ise ilkini alir."""
    d = re.sub(r"\(.*?\)", " ", str(ham or ""))          # parantezli sehir bilgisi
    d = d.replace("*", "").strip(" .,")
    d = re.split(r"\s*(?:/|,|·|;| ve | and )\s*", d)[0].strip()
    return ULKE_ESI.get(sade(d), d)

BASLIK = PatternFill("solid", fgColor="1F2937")
BASLIK_YAZI = Font(bold=True, color="FFFFFF", size=11)
LINK = Font(color="0563C1", underline="single")
VURGU = PatternFill("solid", fgColor="FFF4CE")


def sade(s):
    s = unicodedata.normalize("NFKC", str(s or "")).strip().lower()
    return re.sub(r"\s+", " ", s.replace("*", "").replace("`", ""))


# Bu baslik gorulunce dosyanin geri kalani okunmaz — oradan sonrasi
# ELENEN firmalar. Onlari hedef listesine karistirmak listeyi coper.
DUR = re.compile(r"^#{1,6}\s*.*(elenen|reddedilen|rejected|excluded|dışarıda|disarida)",
                 re.I)


def markdown_tablosu(metin):
    """Dosyadaki TUM markdown tablolarini birlestirip (basliklar, satirlar) dondurur.

    Ajanlar tabloyu bazen ulke ulke boluyor ("## Hindistan (9 firma)" + tablo,
    "## Almanya (1 firma)" + tablo...). Tek tablo okumak sessizce veri kaybettirir:
    33 firmalik dosyadan 9 firma cikmisti. O yuzden hepsi taranir.
    "Elenenler" baslinda tarama durur.
    """
    satirlar = [s.rstrip() for s in metin.split("\n")]

    def hucreler(s):
        return [h.strip() for h in s.strip().strip("|").split("|")]

    def gereksiz(b):
        return sade(b) in ("#", "no", "sıra", "sira", "nr", "")

    basliklar, veri = [], []
    i = 0
    while i < len(satirlar):
        s = satirlar[i]
        if DUR.match(s.strip()):
            break
        ayracli = i + 1 < len(satirlar) and re.match(r"^\s*\|[\s:|-]+\|\s*$", satirlar[i + 1])
        # Ajan ayrac satirini unutursa tablo hic taninmiyordu: 5 bolge dosyasi, 162 firma
        # sessizce kayboluyordu. "| Firma |" ile baslayan satir ayracsiz da baslik sayilir.
        ayracsiz = re.match(r"^\s*\|\s*(#\s*\|\s*)?(firma|company)\s*\|", s, re.I)
        if s.lstrip().startswith("|") and (ayracli or ayracsiz):
            ham = hucreler(s)
            tut = [k for k, b in enumerate(ham) if not gereksiz(b)]
            bu = [ham[k] for k in tut]
            if not basliklar:
                basliklar = bu
            j = i + 2 if ayracli else i + 1
            while j < len(satirlar) and satirlar[j].lstrip().startswith("|"):
                h = hucreler(satirlar[j])
                satir = [(h[k] if k < len(h) else "") for k in tut]
                if any(x for x in satir):
                    # Basliklari ayni olmayan tabloyu kendi sirasina gore hizala
                    if bu != basliklar:
                        eslesme = {sade(b): d for b, d in zip(bu, satir)}
                        satir = [eslesme.get(sade(b), "") for b in basliklar]
                    veri.append((satir + [""] * len(basliklar))[: len(basliklar)])
                j += 1
            i = j
            continue
        i += 1
    return basliklar, veri


def esle(basliklar):
    """Gelen basliklari kanonik semaya oturtur. Doner: (kanonik_basliklar, indeksler)."""
    kalan = list(range(len(basliklar)))
    indeks = {}
    for kanonik, kalip in SEMA:
        for i in list(kalan):
            if re.search(kalip, sade(basliklar[i])):
                indeks[kanonik] = i
                kalan.remove(i)
                break
    cikti = [k for k, _ in SEMA]
    ek = [basliklar[i] for i in kalan]
    return cikti + ek, [indeks.get(k) for k, _ in SEMA] + kalan


def sayfa_yaz(ws, basliklar, satirlar, segment_sutunu=False):
    tum = basliklar + (["Segment"] if segment_sutunu else []) + TAKIP
    ws.append(tum)
    for h in ws[1]:
        h.fill, h.font = BASLIK, BASLIK_YAZI
        h.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 24

    baglanti = [i for i, b in enumerate(tum)
                if re.search(r"site|url|sayfa|iletişim|e-posta|link|taslak", sade(b)) and sade(b) != "hazır e-posta"]

    for satir in satirlar:
        ws.append(list(satir) + [""] * (len(tum) - len(satir)))
        n = ws.max_row
        for i in baglanti:
            h = ws.cell(row=n, column=i + 1)
            d = str(h.value or "").strip()
            m = re.search(r"https?://\S+", d)
            hedef = None
            if m:
                hedef = m.group(0).rstrip(").,;")
            elif re.fullmatch(r"[^@\s]+@[^@\s]+\.[A-Za-z]{2,}", d):
                hedef = "mailto:" + d
            if d.startswith("mailto:") and "?subject=" in d:
                h.value, h.hyperlink, h.font = "✉ Taslağı aç", d, LINK
                continue
            if hedef and len(hedef) < 255:
                h.hyperlink, h.font = hedef, LINK
        # Takip sutunlarini goze carpar yap — doldurulacak alan orasi
        for k in range(len(TAKIP)):
            ws.cell(row=n, column=len(tum) - k).fill = VURGU

    for i, b in enumerate(tum, start=1):
        en = max([len(str(b))] + [len(str(ws.cell(row=r, column=i).value or ""))
                                  for r in range(2, ws.max_row + 1)] or [10])
        ws.column_dimensions[get_column_letter(i)].width = min(max(en + 2, 12), 48)
        if sade(b) in ("ne üretiyor", "doğrulama sayfası"):
            ws.column_dimensions[get_column_letter(i)].width = 46

    ws.freeze_panes = "B2"
    if ws.max_row > 1:
        t = Table(displayName="T" + re.sub(r"\W", "", ws.title)[:20] + str(ws.max_row),
                  ref="A1:%s%d" % (get_column_letter(len(tum)), ws.max_row))
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)


def main():
    if not os.path.isdir(KLASOR):
        sys.exit("[HATA] klasör yok: " + KLASOR)
    dosyalar = sorted(f for f in os.listdir(KLASOR) if f.endswith(".md"))
    if not dosyalar:
        sys.exit("[HATA] seo/hedef-firmalar içinde .md yok — araştırma henüz bitmemiş.")

    wb = Workbook()
    wb.remove(wb.active)
    ozet = wb.create_sheet("TÜM FİRMALAR")
    kanonik = [k for k, _ in SEMA] + EK_BASLIK
    hepsi, kayit, gorulen = [], [], {}

    # --- 1. asama: tum dosyalardan satirlari SEGMENTINE gore topla ---
    # bolge-*.md dosyalarinda her satir kendi segmentini "Segment" sutununda
    # tasir (bolge ajani 14 segmentin hepsini arar); digerlerinde dosya = segment.
    n = len(SEMA)
    toplanan = {}        # segment -> {"basliklar": [...], "satirlar": [...], "bolge": int}
    for dosya in dosyalar:
        anahtar = dosya[:-3]
        basliklar, satirlar = markdown_tablosu(
            io.open(os.path.join(KLASOR, dosya), encoding="utf-8").read())
        if not satirlar:
            print("  [atlandı] %s — tablo bulunamadı" % dosya)
            continue
        bolge = anahtar.startswith("bolge-")
        seg_sutun = next((i for i, b in enumerate(basliklar) if sade(b) == "segment"), None)
        if bolge and seg_sutun is None:
            print("  [atlandı] %s — bölge dosyasında Segment sütunu yok" % dosya)
            continue
        segler_satir = []
        if seg_sutun is not None:
            segler_satir = [segment_bul(s[seg_sutun] if seg_sutun < len(s) else "") for s in satirlar]
            basliklar = basliklar[:seg_sutun] + basliklar[seg_sutun + 1:]
            satirlar = [s[:seg_sutun] + s[seg_sutun + 1:] for s in satirlar]
        yeni_basliklar, indeksler = esle(basliklar)
        duzgun = [[(satir[i] if i is not None and i < len(satir) else "") for i in indeksler]
                  for satir in satirlar]

        # Ulke sutununu tek yazima cek; icindeki sehir bilgisini "Ne üretiyor"a tasi
        # ki suzgec calissin ama bilgi kaybolmasin.
        for satir in duzgun:
            ham = str(satir[1] or "")
            temiz = ulke_duzelt(ham)
            if temiz != ham.strip():
                artik = re.sub(r"[*]", "", ham).replace(temiz, "").strip(" ()/,.·;")
                if artik and sade(artik) not in sade(satir[4]):
                    satir[4] = (artik + " — " + str(satir[4] or "")).strip(" —")
            satir[1] = temiz

        if bolge:
            taninmayan = 0
            for satir, seg in zip(duzgun, segler_satir):
                if not seg:
                    taninmayan += 1
                    continue
                t = toplanan.setdefault(seg, {"basliklar": [k for k, _ in SEMA], "satirlar": [], "bolge": 0})
                t["satirlar"].append(satir[:n] + [""] * (len(t["basliklar"]) - n))
                t["bolge"] += 1
            if taninmayan:
                print("  [uyarı] %s — %d satırın segmenti tanınmadı, atlandı" % (dosya, taninmayan))
        else:
            ad = ADLAR.get(anahtar, anahtar.replace("-", " ").title())
            t = toplanan.setdefault(ad, {"basliklar": yeni_basliklar, "satirlar": [], "bolge": 0})
            if len(t["basliklar"]) < len(yeni_basliklar):
                t["basliklar"] = yeni_basliklar
            t["satirlar"] = duzgun + t["satirlar"]      # segment dosyasi once, bolge satirlari sonra

    # --- 2. asama: segment sayfalarini yaz (segment icinde de tekillestirerek) ---
    sira = {ad: i for i, ad in enumerate(sorted(ADLAR.values(), key=lambda a: [k for k, v in ADLAR.items() if v == a][0]))}
    for ad in sorted(toplanan, key=lambda a: sira.get(a, 999)):
        t = toplanan[ad]
        gorulen_seg, satirlar = set(), []
        for satir in t["satirlar"]:
            k = firma_anahtari(satir)
            if k in gorulen_seg:
                continue
            gorulen_seg.add(k)
            satirlar.append(satir + [""] * (len(t["basliklar"]) - len(satir)))
        yeni_basliklar = t["basliklar"][:n] + EK_BASLIK + t["basliklar"][n:]
        for satir in satirlar:
            satir[n:n] = ek_sutunlar([ad], satir)
        ws_seg = wb.create_sheet(ad[:31])
        sayfa_yaz(ws_seg, yeni_basliklar, satirlar)
        durum_isaretle(ws_seg, segment_adi=ad)

        for satir in satirlar:
            anahtar_firma = firma_anahtari(satir)
            if anahtar_firma in gorulen:
                if ad not in gorulen[anahtar_firma]:
                    gorulen[anahtar_firma].append(ad)
                continue
            gorulen[anahtar_firma] = [ad]
            hepsi.append((anahtar_firma, list(satir[:len(SEMA)])))
        kayit.append((ad, len(satirlar)))
        ek_bilgi = " (%d'i bölge taramasından)" % t["bolge"] if t["bolge"] else ""
        print("  %s: %d firma%s" % (ad, len(satirlar), ek_bilgi))

    if hepsi:
        # Iki segmentte cikan firma tek satir olur ama segmentlerin HEPSI yazilir:
        # kablo kanali + raf ureten firmaya iki hat birden teklif edilir.
        birlesik = []
        for k, satir in hepsi:
            segler = sorted(gorulen[k], key=lambda x: (x in TEYITSIZ, x in SONA))
            birlesik.append(satir + ek_sutunlar(segler, satir) + [" + ".join(segler)])
        sayfa_yaz(ozet, kanonik, birlesik, segment_sutunu=True)
        bekleyen = durum_isaretle(ozet)
        print("  %d firma BEKLE ile işaretlendi (yalnızca teyitsiz segmentte)" % bekleyen)
    else:
        ozet.append(["Hiçbir dosyada tablo bulunamadı."])

    wb.move_sheet("TÜM FİRMALAR", offset=-len(wb.sheetnames) + 1)
    wb.save(CIKTI)

    cift = sum(1 for v in gorulen.values() if len(v) > 1)
    print("\n" + CIKTI)
    print("Toplam %d benzersiz firma, %d segment — %s" % (len(hepsi), len(kayit), date.today()))
    if cift:
        print("%d firma birden fazla segmentte çıktı, TÜM FİRMALAR sayfasında bir kez yazıldı." % cift)


if __name__ == "__main__":
    main()
